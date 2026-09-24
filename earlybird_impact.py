import openpyxl, json
from collections import defaultdict

src = open('analyze_0830.py').read()
start = src.index("ROSTER = {")
end = src.index("all_rows = []")
exec(compile(src[start:end], 'roster', 'exec'))

files = [('0805','redtop2000_0805.xlsx'),('0817','redtop2000merchantsstatus-0817.xlsx'),
         ('0819','redtop2000merchantsstatus-0819.xlsx'),('0820','redtop2000merchantsstatus-0820.xlsx'),
         ('0823','redtop2000merchantsstatus-0823.xlsx'),('0825','redtop2000merchantsstatus-0825.xlsx'),
         ('0827','redtop2000merchantsstatus-0827.xlsx'),('0830','redtop2000merchantsstatus-0830.xlsx')]

hist = defaultdict(list)
for d, f in files:
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.worksheets[0]
    hdr = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
    oph = [h for h in hdr if 'operating' in h][0]
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(hdr, row))
        hist[str(r['Lead ID'])].append((d, r.get(oph), r.get('BD')))

wb = openpyxl.load_workbook('redtop2000merchantsstatus-0830.xlsx', read_only=True, data_only=True)
ws = wb['Sheet1']
hdr = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
oph = '0824-0830 operating'
counted = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr, row))
    if r['BD'] in all_roster_bds and r.get('target type') in ('not sign','not online','not operating') and r.get(oph) == 1:
        counted.append(r)

# early bird eligibility per counted lead
eb_solid, eb_pending = defaultdict(int), defaultdict(int)
for r in counted:
    lid = str(r['Lead ID'])
    h = hist.get(lid, [])
    first_op = min((d for d, op, bd in h if op == 1), default=None)
    if first_op == '0805':
        eb_solid[r['BD']] += 1
    elif first_op == '0817':
        eb_pending[r['BD']] += 1

D = json.load(open('order_penetration_data_0830_final.json'))

def scenario(extra):
    newtot = {}
    for b in D['ws_bds'] + D['santos_bds']:
        newtot[b['bd']] = b['total'] + eb_solid.get(b['bd'], 0) + (eb_pending.get(b['bd'], 0) if extra else 0)
    return newtot

print('=== Scenario A: 仅铁证早鸟 (0805文件已op=1, 43条) ===')
na = scenario(False)
qs_a = sum(1 for v in na.values() if v >= 15)
for b in sorted(D['ws_bds'] + D['santos_bds'], key=lambda x: -na[x['bd']]):
    delta = na[b['bd']] - b['total']
    if delta or b['total'] >= 12:
        flag = '→达标' if (b['total'] < 15 <= na[b['bd']]) else ('掉出达标' if (b['total'] >= 15 > na[b['bd']]) else '')
        print(f"  {b['bd']:18} {b['total']:3} → {na[b['bd']]:3} (+{delta}) {flag} [{b['bdm']}]")

print('\n=== Scenario B: 铁证+待验证全部计入 (43+52=95条) ===')
nb = scenario(True)
for b in sorted(D['ws_bds'] + D['santos_bds'], key=lambda x: -nb[x['bd']]):
    delta = nb[b['bd']] - b['total']
    if b['total'] >= 12 or delta >= 2:
        flag = '→达标' if (b['total'] < 15 <= nb[b['bd']]) else ''
        print(f"  {b['bd']:18} {b['total']:3} → {nb[b['bd']]:3} (+{delta}) {flag} [{b['bdm']}]")

q0 = sum(1 for b in D['ws_bds']+D['santos_bds'] if b['total'] >= 15)
qa = sum(1 for v in na.values() if v >= 15)
qb = sum(1 for v in nb.values() if v >= 15)
print(f'\n达标BD: 现在{q0} → A方案{qa} → B方案{qb}')

print('\n=== BDM 人均 (A方案) ===')
for bdm, v in sorted(D['bdm_summary'].items(), key=lambda x: -x[1]['avg']):
    team = v['team_score'] + sum(eb_solid.get(bd, 0) for bd in bd_to_bdm if bd_to_bdm[bd] == bdm)
    team_b = v['team_score'] + sum(eb_solid.get(bd, 0) + eb_pending.get(bd, 0) for bd in bd_to_bdm if bd_to_bdm[bd] == bdm)
    avg_a = team / v['bd_count']
    avg_b = team_b / v['bd_count']
    print(f"  {bdm:22} avg {v['avg']:6} → A:{avg_a:6.1f}{'✅' if avg_a>=15 else ' '} B:{avg_b:6.1f}{'✅' if avg_b>=15 else ' '} [{v['city']}]")

print('\n=== Adriana 7案逐条核验 ===')
cases = {'50198082':('tiagodangelo','Nosso Cafés'),'50203029':('tiagodangelo','Padaria Cristal'),
         '50175552':('marciajesuino','Niku Burger'),'50176441':('marciajesuino','Sailor Dog'),
         '50186043':('wanessasilva','Casa de Bolos'),'50175774':('wanessasilva','A Bem Dita'),'50195890':('wanessasilva','Ponto da Esfiha')}
for lid, (bd, name) in cases.items():
    h = hist.get(lid, [])
    first_op = min((d for d, op, b in h if op == 1), default=None)
    bd0805 = next((b for d, op, b in h if d == '0805' and op == 1), None)
    verdict = '✅ 早鸟成立(0805已营业)' if first_op == '0805' else ('⚠️ 首次营业在0817窗口(0806-0817), 需证据证明≤0815' if first_op == '0817' else '❌ 无早鸟')
    extra = f' (0805时BD={bd0805})' if bd0805 and bd0805 != bd else ''
    print(f'  {name:16} ({bd:14}): first_op={first_op} {verdict}{extra}')
