import openpyxl
import json
from collections import defaultdict

wb = openpyxl.load_workbook('redtop2000merchantsstatus-0819.xlsx', data_only=True)
ws = wb['Sheet1']
headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
cols = {h: i+1 for i, h in enumerate(headers) if h}

print("=== Excel Headers ===")
for h in headers:
    print(f"  Col {cols.get(h,'?')}: {h}")

metro_cities = {'Southern São Paulo Metropolitan', 'Western São Paulo Metropolitan', 'Santos City'}

ROSTER = {
    'Southern São Paulo Metropolitan': {
        'adriananaves': ['brunacrelien', 'elisangelasouza', 'evaldosilva', 'ivanfelizardo', 'josenascimento', 'murilosilva', 'rodrigoalmeida', 'tiagodangelo', 'wanessasilva'],
        'fernandooliveira': ['brunapadilha', 'carlosmotta', 'erikboilesen', 'felipesanches', 'pedrosaccone', 'wainnergonzales', 'wellingtonsantos'],
        'eduardoalbuquerque': ['allanmaalouli', 'cristianedasilva', 'felipesilva', 'fernandoquintiliano', 'joycepurificacao', 'otaviomazzega', 'thaisvasconcelos'],
        'alisaeed': ['camilabatista', 'davidcaramaschi', 'eduardosantos', 'marciajesuino', 'ricardoaraujo', 'rodrigocorreia'],
        'tadeumoraes': ['leandroresende', 'lucasreis', 'matheuscoelho', 'paulosantos', 'paulosilva', 'ricardoalmeida'],
    },
    'Western São Paulo Metropolitan': {
        'biancaceotto': ['caiolorencato', 'karinacolomina', 'luizmoreira', 'nayannedias', 'nicolemenezes', 'nilmaxfranca', 'rodrigopoli'],
        'cesararraes': ['guilhermesantos', 'jadycarvalho', 'johnataguimaraes', 'kevinberti', 'renansilva', 'williamhenrique'],
        'thiagoscavazini': ['barbaraoliveira', 'claudiareis', 'felipecruz', 'heloisamagnani', 'marialisboa', 'milenasantana'],
        'igorfeitosa': ['alexlima', 'igorsouza', 'marcelofeitosa', 'tamirislopes', 'tatianarodrigues'],
        'lucasferreira': ['clebersoares', 'lucasnada', 'tatianeribeiro', 'valeriachacon'],
    },
    'Santos City': {
        'sabrinafernandes': ['brunosalmaso', 'ericabarbosa', 'felipebarbosa', 'giovanabareno'],
        'renataleite': ['jessicarossi', 'joaocastro', 'luizmarques', 'nathaliabernardino'],
    },
}

all_roster_bds = set()
bd_to_bdm = {}
bd_to_city = {}
for city_full, bdm_map in ROSTER.items():
    city_short = city_full.split(' São Paulo')[0] if 'São Paulo' in city_full else city_full.replace(' City','')
    for bdm, bd_list in bdm_map.items():
        for bd in bd_list:
            all_roster_bds.add(bd)
            bd_to_bdm[bd] = bdm
            bd_to_city[bd] = city_short

# Collect ALL metro rows
all_rows = []
bd_seen = set()
for r in range(2, ws.max_row+1):
    region = ws.cell(r, cols['region']).value
    city = ws.cell(r, cols['city']).value
    if region == 'Metropolitan Region' and city in metro_cities:
        row = {
            'lead_name': ws.cell(r, cols['Lead Name']).value,
            'red_order': ws.cell(r, cols['red order']).value,
            'city': city,
            'target_type': ws.cell(r, cols['target type-删除offline，删除近30日签约']).value,
            'bd': ws.cell(r, cols['BD']).value,
            'bdm': ws.cell(r, cols['BDM']).value,
            'sign': ws.cell(r, cols['sign-0819']).value,
            'operating': ws.cell(r, cols['0813-0819 operating']).value,
        }
        all_rows.append(row)

print(f"\nTotal metro rows: {len(all_rows)}")
print(f"Roster BDs: {len(all_roster_bds)}")

# Direction mapping
dir_map = {
    'not sign': ('not_signed', 6),
    'not online': ('not_online', 4),
    'not operating': ('not_operating', 3),
}

# ===== REGIONAL OVERVIEW (ALL merchants) =====
total_target = len(all_rows)
total_signed = sum(1 for r in all_rows if r['sign'] == 1)
total_operating = sum(1 for r in all_rows if r['operating'] == 1)

ns_target_all = sum(1 for r in all_rows if r['target_type'] == 'not sign')
no_target_all = sum(1 for r in all_rows if r['target_type'] == 'not online')
nop_target_all = sum(1 for r in all_rows if r['target_type'] == 'not operating')

ns_op_all = sum(1 for r in all_rows if r['target_type'] == 'not sign' and r['operating'] == 1)
no_op_all = sum(1 for r in all_rows if r['target_type'] == 'not online' and r['operating'] == 1)
nop_op_all = sum(1 for r in all_rows if r['target_type'] == 'not operating' and r['operating'] == 1)

# Total score (ALL operating merchants)
total_score_all = 0
hv_count_all = 0
ns_score_all = 0
no_score_all = 0
nop_score_all = 0
hv_bonus_all = 0
for r in all_rows:
    if r['target_type'] not in dir_map:
        continue
    if r['operating'] != 1:
        continue
    direction, base = dir_map[r['target_type']]
    hv = 1 if (r['red_order'] and r['red_order'] >= 20) else 0
    score = base + (2 if hv else 0)
    total_score_all += score
    if direction == 'not_signed': ns_score_all += score
    elif direction == 'not_online': no_score_all += score
    elif direction == 'not_operating': nop_score_all += score
    if hv:
        hv_count_all += 1
        hv_bonus_all += 2

# BD-level score (roster BDs only)
bd_scores = defaultdict(lambda: {'not_signed': 0, 'not_online': 0, 'not_operating': 0, 'hv_bonus': 0, 'total': 0, 'leads': []})
for r in all_rows:
    bd = r['bd']
    if bd not in all_roster_bds:
        continue
    tt = r['target_type']
    if tt not in dir_map:
        continue
    if r['operating'] != 1:
        continue
    direction, base = dir_map[tt]
    hv = 1 if (r['red_order'] and r['red_order'] >= 20) else 0
    score = base + (2 if hv else 0)
    bd_scores[bd][direction] += 1
    bd_scores[bd]['hv_bonus'] += (2 if hv else 0)
    bd_scores[bd]['total'] += score
    bd_scores[bd]['leads'].append({
        'name': r['lead_name'],
        'direction': direction,
        'base': base,
        'hv': hv,
        'score': score,
    })

bd_level_score = sum(sc['total'] for sc in bd_scores.values())

print("\n===== REGIONAL OVERVIEW (ALL) =====")
print(f"Target: {total_target}")
print(f"Signed: {total_signed} ({total_signed/total_target*100:.1f}%)")
print(f"Operating: {total_operating} ({total_operating/total_target*100:.1f}%)")
print(f"Total Score: {total_score_all}")
print(f"BD-Level Score: {bd_level_score}")
print(f"NS: target={ns_target_all}, op={ns_op_all}, score={ns_score_all}")
print(f"NO: target={no_target_all}, op={no_op_all}, score={no_score_all}")
print(f"NOP: target={nop_target_all}, op={nop_op_all}, score={nop_score_all}")
print(f"HV: {hv_count_all} merchants, +{hv_bonus_all} pts")

# ===== CITY OVERVIEW =====
print("\n===== CITY OVERVIEW =====")
city_data = {}
for city_full in ['Southern São Paulo Metropolitan', 'Western São Paulo Metropolitan', 'Santos City']:
    city_short = city_full.split(' São Paulo')[0] if 'São Paulo' in city_full else city_full.replace(' City','')
    c_rows = [r for r in all_rows if r['city'] == city_full]
    c_target = len(c_rows)
    c_signed = sum(1 for r in c_rows if r['sign'] == 1)
    c_op = sum(1 for r in c_rows if r['operating'] == 1)
    
    # Assigned = BD is in roster
    c_assigned = sum(1 for r in c_rows if r['bd'] in all_roster_bds)
    
    c_ns_target = sum(1 for r in c_rows if r['target_type'] == 'not sign')
    c_no_target = sum(1 for r in c_rows if r['target_type'] == 'not online')
    c_nop_target = sum(1 for r in c_rows if r['target_type'] == 'not operating')
    c_ns_op = sum(1 for r in c_rows if r['target_type'] == 'not sign' and r['operating'] == 1)
    c_no_op = sum(1 for r in c_rows if r['target_type'] == 'not online' and r['operating'] == 1)
    c_nop_op = sum(1 for r in c_rows if r['target_type'] == 'not operating' and r['operating'] == 1)
    
    # Score (all operating in city)
    c_score = 0
    c_hv = 0
    c_ns_score = 0
    c_no_score = 0
    c_nop_score = 0
    c_hv_bonus = 0
    for r in c_rows:
        if r['target_type'] not in dir_map:
            continue
        if r['operating'] != 1:
            continue
        direction, base = dir_map[r['target_type']]
        hv = 1 if (r['red_order'] and r['red_order'] >= 20) else 0
        score = base + (2 if hv else 0)
        c_score += score
        if direction == 'not_signed': c_ns_score += score
        elif direction == 'not_online': c_no_score += score
        elif direction == 'not_operating': c_nop_score += score
        if hv:
            c_hv += 1
            c_hv_bonus += 2
    
    city_data[city_short] = {
        'target': c_target,
        'assigned': c_assigned,
        'unassigned': c_target - c_assigned,
        'assigned_pct': round(c_assigned/c_target*100, 1) if c_target else 0,
        'signed': c_signed,
        'operating': c_op,
        'score': c_score,
        'op_rate': round(c_op/c_target*100, 1) if c_target else 0,
        'operating_pct': round(c_op/c_target*100, 1) if c_target else 0,
        'ns_target': c_ns_target, 'no_target': c_no_target, 'nop_target': c_nop_target,
        'ns_op': c_ns_op, 'no_op': c_no_op, 'nop_op': c_nop_op,
        'ns_score': c_ns_score, 'no_score': c_no_score, 'nop_score': c_nop_score,
        'hv_count': c_hv, 'hv_bonus': c_hv_bonus,
        'directions': {
            'not sign': {'target': c_ns_target, 'operating': c_ns_op, 'score': c_ns_score,
                         'pct': round(c_ns_op/c_ns_target*100, 1) if c_ns_target else 0},
            'not online': {'target': c_no_target, 'operating': c_no_op, 'score': c_no_score,
                          'pct': round(c_no_op/c_no_target*100, 1) if c_no_target else 0},
            'not operating': {'target': c_nop_target, 'operating': c_nop_op, 'score': c_nop_score,
                             'pct': round(c_nop_op/c_nop_target*100, 1) if c_nop_target else 0},
        }
    }
    print(f"  {city_short}: target={c_target}, assigned={c_assigned}, signed={c_signed}, op={c_op}, score={c_score}")
    print(f"    NS: target={c_ns_target}, op={c_ns_op}, score={c_ns_score}")
    print(f"    NO: target={c_no_target}, op={c_no_op}, score={c_no_score}")
    print(f"    NOP: target={c_nop_target}, op={c_nop_op}, score={c_nop_score}")
    print(f"    HV: {c_hv} merchants, +{c_hv_bonus}")

# ===== BDM SCORES =====
print("\n===== BDM SCORES =====")
bdm_data = {}
for city_full, bdm_map in ROSTER.items():
    city_short = city_full.split(' São Paulo')[0] if 'São Paulo' in city_full else city_full.replace(' City','')
    for bdm, bd_list in bdm_map.items():
        team_score = sum(bd_scores[bd]['total'] for bd in bd_list if bd in bd_scores)
        bd_count = len(bd_list)
        avg = team_score / bd_count if bd_count > 0 else 0
        bdm_data[bdm] = {
            'city': city_short,
            'bd_count': bd_count,
            'team_score': team_score,
            'score': team_score,
            'avg': round(avg, 2),
            'qualified': avg >= 15,
        }

for bdm, d in sorted(bdm_data.items(), key=lambda x: -x[1]['avg']):
    status = "✅" if d['qualified'] else "⚠️" if d['avg'] >= 10 else "❌"
    print(f"  {bdm} (city={d['city']}): BDs={d['bd_count']}, TeamScore={d['team_score']}, Avg={d['avg']} {status}")

# ===== BD SCORES =====
ws_bds = []
santos_bds = []
for bd in all_roster_bds:
    sc = bd_scores.get(bd, {'not_signed': 0, 'not_online': 0, 'not_operating': 0, 'hv_bonus': 0, 'total': 0, 'leads': []})
    city = bd_to_city[bd]
    bdm = bd_to_bdm[bd]
    
    # Build details for dashboard
    details = []
    for lead in sc.get('leads', []):
        dir_label_map = {'not_signed': 'not sign', 'not_online': 'not online', 'not_operating': 'not operating'}
        details.append({
            'name': lead['name'],
            'dir': dir_label_map.get(lead['direction'], lead['direction']),
            'base': lead['base'],
            'hv': lead['hv'],
            'sub': lead['score'],
        })
    
    entry = {
        'bd': bd, 'bdm': bdm, 'city': city,
        'ns': sc['not_signed'], 'no': sc['not_online'], 'nop': sc['not_operating'],
        'hv': sc['hv_bonus'] // 2,  # count of HV merchants
        'total': sc['total'],
        'details': details,
        'not_signed': sc['not_signed'], 'not_online': sc['not_online'], 'not_operating': sc['not_operating'],
        'hv_bonus': sc['hv_bonus'],
        'leads': sc.get('leads', []),
    }
    if city == 'Santos':
        santos_bds.append(entry)
    else:
        ws_bds.append(entry)

ws_bds.sort(key=lambda x: -x['total'])
santos_bds.sort(key=lambda x: -x['total'])

print(f"\n===== WS BDs (top 10) =====")
for i, bd in enumerate(ws_bds[:10]):
    print(f"  {i+1}. {bd['bd']}: Total={bd['total']}, NS={bd['ns']}, NO={bd['no']}, NOP={bd['nop']}, HV={bd['hv']}")

print(f"\n===== Santos BDs (all) =====")
for i, bd in enumerate(santos_bds):
    print(f"  {i+1}. {bd['bd']}: Total={bd['total']}, NS={bd['ns']}, NO={bd['no']}, NOP={bd['nop']}, HV={bd['hv']}")

qualified_bds = [bd for bd in ws_bds + santos_bds if bd['total'] >= 15]
print(f"\nQualified BDs (≥15): {len(qualified_bds)}")

# ===== LOAD PREVIOUS DATA (0817) =====
with open('order_penetration_data_0817_final.json') as f:
    prev_data = json.load(f)

prev = {
    'date': '0817',
    'region': {
        'total': prev_data['regional_summary']['total_target'],
        'assigned': sum(prev_data['city_summary'][c]['assigned'] if 'assigned' in prev_data['city_summary'][c] else 0 for c in prev_data['city_summary']),
        'operating': prev_data['regional_summary']['total_operating'],
        'score': prev_data['regional_summary']['total_score'],
    },
    'directions': {
        'not sign': {'target': prev_data['regional_summary']['ns_target'], 'operating': prev_data['regional_summary']['ns_op'], 'score': prev_data['regional_summary']['ns_score']},
        'not online': {'target': prev_data['regional_summary']['no_target'], 'operating': prev_data['regional_summary']['no_op'], 'score': prev_data['regional_summary']['no_score']},
        'not operating': {'target': prev_data['regional_summary']['nop_target'], 'operating': prev_data['regional_summary']['nop_op'], 'score': prev_data['regional_summary']['nop_score']},
    },
    'cities': {}
}

# Build prev city data with directions
for c in ['Southern', 'Western', 'Santos']:
    pc = prev_data['city_summary'][c]
    prev['cities'][c] = {
        'target': pc['target'],
        'assigned': pc.get('assigned', 0),
        'operating': pc['operating'],
        'score': pc['score'],
        'directions': {
            'not sign': {'target': pc['ns_target'], 'operating': pc['ns_op'], 'score': pc['ns_score'], 'pct': round(pc['ns_op']/pc['ns_target']*100,1) if pc['ns_target'] else 0},
            'not online': {'target': pc['no_target'], 'operating': pc['no_op'], 'score': pc['no_score'], 'pct': round(pc['no_op']/pc['no_target']*100,1) if pc['no_target'] else 0},
            'not operating': {'target': pc['nop_target'], 'operating': pc['nop_op'], 'score': pc['nop_score'], 'pct': round(pc['nop_op']/pc['nop_target']*100,1) if pc['nop_target'] else 0},
        }
    }

# Also need assigned count for current data
total_assigned = sum(city_data[c]['assigned'] for c in city_data)

# ===== BUILD DASHBOARD DATA JSON =====
dashboard_data = {
    'data_date': '0819',
    'operating_period': '0813-0819',
    'region': {
        'total': total_target,
        'assigned': total_assigned,
        'assigned_pct': round(total_assigned/total_target*100, 1) if total_target else 0,
        'operating': total_operating,
        'operating_pct': round(total_operating/total_target*100, 1) if total_target else 0,
        'score': total_score_all,
    },
    'directions': {
        'not sign': {'target': ns_target_all, 'operating': ns_op_all, 'score': ns_score_all,
                     'pct': round(ns_op_all/ns_target_all*100, 1) if ns_target_all else 0},
        'not online': {'target': no_target_all, 'operating': no_op_all, 'score': no_score_all,
                       'pct': round(no_op_all/no_target_all*100, 1) if no_target_all else 0},
        'not operating': {'target': nop_target_all, 'operating': nop_op_all, 'score': nop_score_all,
                          'pct': round(nop_op_all/nop_target_all*100, 1) if nop_target_all else 0},
    },
    'cities': {},
    'ws_bds': ws_bds,
    'santos_bds': santos_bds,
    'ws_bdms': [],
    'santos_bdms': [],
}

# Build city data for dashboard
for c in ['Southern', 'Western', 'Santos']:
    cd = city_data[c]
    dashboard_data['cities'][c] = {
        'target': cd['target'],
        'assigned': cd['assigned'],
        'assigned_pct': cd['assigned_pct'],
        'unassigned': cd['unassigned'],
        'operating': cd['operating'],
        'operating_pct': cd['operating_pct'],
        'score': cd['score'],
        'directions': cd['directions'],
    }

# Build BDM lists for dashboard (WS = Southern + Western, Santos = Santos)
ws_bdms = []
santos_bdms = []
for bdm, d in bdm_data.items():
    entry = {
        'bdm': bdm,
        'city': d['city'],
        'bd_count': d['bd_count'],
        'score': d['team_score'],
        'avg': d['avg'],
    }
    if d['city'] == 'Santos':
        santos_bdms.append(entry)
    else:
        ws_bdms.append(entry)

ws_bdms.sort(key=lambda x: -x['avg'])
santos_bdms.sort(key=lambda x: -x['avg'])

dashboard_data['ws_bdms'] = ws_bdms
dashboard_data['santos_bdms'] = santos_bdms

# Save dashboard data
with open('/tmp/dashboard_data.json', 'w') as f:
    json.dump(dashboard_data, f, indent=2, default=str)

# Save final data JSON
output = {
    'data_date': '0819',
    'operating_period': '0813-0819',
    'previous': {
        'date': '0817',
        'total_operating': prev_data['regional_summary']['total_operating'],
        'total_score': prev_data['regional_summary']['total_score'],
        'bd_level_score': prev_data['regional_summary']['bd_level_score'],
        'ns_op': prev_data['regional_summary']['ns_op'],
        'no_op': prev_data['regional_summary']['no_op'],
        'nop_op': prev_data['regional_summary']['nop_op'],
        'signed': prev_data['regional_summary']['total_signed'],
        'cities': {
            'Southern': {'signed': prev_data['city_summary']['Southern']['signed'], 'op': prev_data['city_summary']['Southern']['operating'], 'score': prev_data['city_summary']['Southern']['score'], 'ns_op': prev_data['city_summary']['Southern']['ns_op'], 'no_op': prev_data['city_summary']['Southern']['no_op'], 'nop_op': prev_data['city_summary']['Southern']['nop_op']},
            'Western': {'signed': prev_data['city_summary']['Western']['signed'], 'op': prev_data['city_summary']['Western']['operating'], 'score': prev_data['city_summary']['Western']['score'], 'ns_op': prev_data['city_summary']['Western']['ns_op'], 'no_op': prev_data['city_summary']['Western']['no_op'], 'nop_op': prev_data['city_summary']['Western']['nop_op']},
            'Santos': {'signed': prev_data['city_summary']['Santos']['signed'], 'op': prev_data['city_summary']['Santos']['operating'], 'score': prev_data['city_summary']['Santos']['score'], 'ns_op': prev_data['city_summary']['Santos']['ns_op'], 'no_op': prev_data['city_summary']['Santos']['no_op'], 'nop_op': prev_data['city_summary']['Santos']['nop_op']},
        }
    },
    'regional_summary': {
        'total_target': total_target,
        'total_signed': total_signed,
        'total_operating': total_operating,
        'total_score': total_score_all,
        'bd_level_score': bd_level_score,
        'ns_target': ns_target_all, 'no_target': no_target_all, 'nop_target': nop_target_all,
        'ns_op': ns_op_all, 'no_op': no_op_all, 'nop_op': nop_op_all,
        'ns_score': ns_score_all, 'no_score': no_score_all, 'nop_score': nop_score_all,
        'hv_count': hv_count_all, 'hv_bonus': hv_bonus_all,
    },
    'city_summary': city_data,
    'bdm_summary': bdm_data,
    'ws_bds': ws_bds,
    'santos_bds': santos_bds,
    'bd_scores_detail': {bd: {'leads': sc.get('leads', []), 'not_signed': sc['not_signed'], 'not_online': sc['not_online'], 'not_operating': sc['not_operating'], 'hv_bonus': sc['hv_bonus'], 'total': sc['total']} for bd, sc in bd_scores.items()},
}

with open('order_penetration_data_0819_final.json', 'w') as f:
    json.dump(output, f, indent=2, default=str)

print("\nSaved to /tmp/dashboard_data.json")
print("Saved to order_penetration_data_0819_final.json")
print("\n===== DONE =====")
