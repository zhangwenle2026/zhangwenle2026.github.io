# -*- coding: utf-8 -*-
"""在 analyze_0830.py 基础上加早鸟加分(铁证43家: 0805文件已op=1)"""
import openpyxl, json
from collections import defaultdict

src = open('analyze_0830.py').read()
start = src.index("ROSTER = {")
end = src.index("all_rows = []")
exec(compile(src[start:end], 'roster', 'exec'))

# 1. build history from 0805 file: lead -> op
wb0 = openpyxl.load_workbook('redtop2000_0805.xlsx', read_only=True, data_only=True)
ws0 = wb0.worksheets[0]
hdr0 = [str(c.value) if c.value else '' for c in next(ws0.iter_rows(min_row=1, max_row=1))]
oph0 = [h for h in hdr0 if 'operating' in h][0]
early_leads = set()  # lead ids with op=1 in 0805 snapshot
for row in ws0.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr0, row))
    if r.get(oph0) == 1:
        early_leads.add(str(r['Lead ID']))
print(f'0805 snapshot operating leads: {len(early_leads)}')

# 2. load current data json
D = json.load(open('order_penetration_data_0830_final.json'))

# 3. recompute per-BD with early bird: need per-lead recompute from 0830 xlsx
wb = openpyxl.load_workbook('redtop2000merchantsstatus-0830.xlsx', read_only=True, data_only=True)
ws = wb['Sheet1']
hdr = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
oph = '0824-0830 operating'
dir_map = {'not sign': ('not_signed', 6), 'not online': ('not_online', 4), 'not operating': ('not_operating', 3)}

bd_scores = defaultdict(lambda: {'not_signed': 0, 'not_online': 0, 'not_operating': 0, 'hv_bonus': 0, 'hv_count': 0, 'early': 0, 'total': 0, 'leads': []})
seen = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr, row))
    bd = r['BD']
    if bd not in all_roster_bds:
        continue
    tt = r.get('target type')
    if tt not in dir_map or r.get(oph) != 1:
        continue
    lid = str(r['Lead ID'])
    if lid in seen:
        continue
    seen.add(lid)
    direction, base = dir_map[tt]
    hv = 1 if (r['red order'] and r['red order'] >= 20) else 0
    eb = 1 if lid in early_leads else 0
    score = base + (2 if hv else 0) + eb
    bd_scores[bd][direction] += 1
    bd_scores[bd]['hv_bonus'] += (2 if hv else 0)
    bd_scores[bd]['hv_count'] += hv
    bd_scores[bd]['early'] += eb
    bd_scores[bd]['total'] += score
    bd_scores[bd]['leads'].append({
        'name': r['Lead Name'], 'direction': direction, 'base': base,
        'hv': hv, 'early': eb, 'score': score,
    })

total_early = sum(s['early'] for s in bd_scores.values())
print(f'early bird applied: {total_early} leads')

# 4. BDM
bdm_data = {}
for city_full, bdm_map in ROSTER.items():
    city_short = city_full.split(' São Paulo')[0] if 'São Paulo' in city_full else city_full.replace(' City','')
    for bdm, bd_list in bdm_map.items():
        team_score = sum(bd_scores[bd]['total'] for bd in bd_list if bd in bd_scores)
        avg = team_score / len(bd_list)
        bdm_data[bdm] = {'city': city_short, 'bd_count': len(bd_list), 'team_score': team_score,
                         'avg': round(avg, 2), 'qualified': avg >= 15}

# 5. BD lists
ws_bds, santos_bds = [], []
for bd in all_roster_bds:
    sc = bd_scores.get(bd, {'not_signed': 0, 'not_online': 0, 'not_operating': 0, 'hv_bonus': 0, 'hv_count': 0, 'early': 0, 'total': 0, 'leads': []})
    entry = {'bd': bd, 'bdm': bd_to_bdm[bd], 'city': bd_to_city[bd], **sc}
    (santos_bds if bd_to_city[bd] == 'Santos' else ws_bds).append(entry)
ws_bds.sort(key=lambda x: -x['total'])
santos_bds.sort(key=lambda x: -x['total'])

# 6. regional summary: add early bonus to scores
R = D['regional_summary']
R['total_score'] += total_early
R['bd_level_score'] += total_early
R['early_count'] = total_early
R['early_bonus'] = total_early
for city, v in D['city_summary'].items():
    ce = sum(s['early'] for bd, s in bd_scores.items() if bd_to_city[bd] == city)
    v['score'] += ce
    v['early_count'] = ce

# 7. save
D['ws_bds'] = ws_bds
D['santos_bds'] = santos_bds
D['bdm_summary'] = bdm_data
D['bd_scores_detail'] = {bd: {'leads': s['leads'], 'not_signed': s['not_signed'], 'not_online': s['not_online'],
                              'not_operating': s['not_operating'], 'hv_bonus': s['hv_bonus'], 'hv_count': s['hv_count'],
                              'early': s['early'], 'total': s['total']} for bd, s in bd_scores.items()}
D['early_bird_rule'] = '铁证口径: 0805快照已营业 → +1/家; 0817窗口52家待证据后补'
json.dump(D, open('order_penetration_data_0830_final.json', 'w'), indent=2, default=str, ensure_ascii=False)
print('saved order_penetration_data_0830_final.json (early bird applied)')

q = [b for b in ws_bds + santos_bds if b['total'] >= 15]
print(f"Qualified BDs: {len(q)}/71")
for b in ws_bds[:5] + santos_bds[:3]:
    print(f"  {b['bd']:18} total={b['total']} (early=+{b['early']})")
print('Qualified BDMs:', sum(1 for v in bdm_data.values() if v['qualified']))
for bdm, v in sorted(bdm_data.items(), key=lambda x: -x[1]['avg'])[:3]:
    print(f"  {bdm}: avg={v['avg']}")
