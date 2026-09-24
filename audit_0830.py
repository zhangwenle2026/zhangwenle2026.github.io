import openpyxl, json
from collections import defaultdict, Counter

src = open('analyze_0830.py').read()
start = src.index("ROSTER = {")
end = src.index("all_rows = []")
exec(compile(src[start:end], 'roster', 'exec'))

files = [('0806','redtop2000_0805.xlsx'),('0817','redtop2000merchantsstatus-0817.xlsx'),
         ('0819','redtop2000merchantsstatus-0819.xlsx'),('0820','redtop2000merchantsstatus-0820.xlsx'),
         ('0823','redtop2000merchantsstatus-0823.xlsx'),('0825','redtop2000merchantsstatus-0825.xlsx'),
         ('0827','redtop2000merchantsstatus-0827.xlsx'),('0830','redtop2000merchantsstatus-0830.xlsx')]

hist = defaultdict(list)
for d, f in files:
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.worksheets[0]
    hdr = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
    sg = [h for h in hdr if h.startswith('sign-')][0]
    oph = [h for h in hdr if 'operating' in h][0]
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(hdr, row))
        hist[str(r['Lead ID'])].append((d, r.get(oph), r.get('BD'), r.get('BDM'), r.get('target type'), r.get(sg)))

D = json.load(open('order_penetration_data_0830_final.json'))
wb = openpyxl.load_workbook('redtop2000merchantsstatus-0830.xlsx', read_only=True, data_only=True)
ws = wb['Sheet1']
hdr = [str(c.value) if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
oph = '0824-0830 operating'
counted = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr, row))
    if r['BD'] in all_roster_bds and r.get('target type') in ('not sign','not online','not operating') and r.get(oph) == 1:
        counted.append(r)

c = Counter(str(r['Lead ID']) for r in counted)
dups = {k: v for k, v in c.items() if v > 1}
print('=== 1. 重复计分 (same Lead ID counted >1x) ===')
for k in dups:
    for r in counted:
        if str(r['Lead ID']) == k:
            print(' ', k, str(r['Lead Name'])[:28], '| type:', r['target type'], '| BD:', r['BD'])

print('\n=== 2. 早鸟资格扫描 (当前计分leads) ===')
eb_auto, eb_ambig, eb_no = [], [], []
seen = set()
for r in counted:
    lid = str(r['Lead ID'])
    if lid in seen: continue
    seen.add(lid)
    h = hist.get(lid, [])
    op0806 = next((op for d, op, *_ in h if d == '0806'), None)
    ops = [d for d, op, bd, bdm, tt, sg_ in h if op == 1]
    first = min(ops) if ops else None
    if op0806 == 1:
        eb_auto.append((r['BD'], lid, str(r['Lead Name'])[:26], r['target type']))
    elif first and first <= '0817':
        eb_ambig.append((r['BD'], lid, str(r['Lead Name'])[:26], r['target type'], first))
    else:
        eb_no.append((r['BD'], lid, str(r['Lead Name'])[:26], first))
print(f'A. 0806文件已op=1 → 早鸟铁证: {len(eb_auto)}条')
by_bd = defaultdict(int)
for bd, *_ in eb_auto: by_bd[bd] += 1
for bd, n in sorted(by_bd.items(), key=lambda x: -x[1]): print(f'   {bd}: +{n}')
print(f'\nB. 首次op=1在0817 (窗口0807-0817, 无法证实≤0815): {len(eb_ambig)}条')
for bd, lid, name, tt, first in sorted(eb_ambig): print(f'   {bd:16} {lid} {name:26} {tt} first_op={first}')
print(f'\nC. 0819后才转化 (无早鸟): {len(eb_no)}条')

print('\n=== 3. Casa de Lanches (Lead 50469429) 完整历史 ===')
for d, op, bd, bdm, tt, sg_ in sorted(hist.get('50469429', [])):
    print(f'   {d}: op={op} sign={sg_} BD={bd} BDM={bdm} type={tt}')

print('\n=== 4. marciajesuino 归属历史 ===')
mb = defaultdict(set)
for lid, rows in hist.items():
    for d, op, bd, bdm, tt, sg_ in rows:
        if bd == 'marciajesuino':
            mb[d].add(bdm)
for d in sorted(mb): print(f'   {d}: {mb[d]}')

print('\n=== 5. 0830文件中BD与roster BDM不一致的行 (counted) ===')
mism = defaultdict(list)
for r in counted:
    if bd_to_bdm.get(r['BD']) != r.get('BDM'):
        mism[(r['BD'], r.get('BDM'), bd_to_bdm.get(r['BD']))].append(str(r['Lead Name'])[:24])
for k, v in sorted(mism.items()):
    print(f'   BD={k[0]} excel_BDM={k[1]} roster_BDM={k[2]}: {len(v)}条 -> {v[:4]}')
